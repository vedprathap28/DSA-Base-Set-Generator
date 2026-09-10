"""
xlsx_generator.py
Converts parsed JSON rows into a platform-ready .xlsx file.
Enforces exact 14-column spec, Arial 10, wrap text, frozen header.
Sheet tab: BASE_QUE_SET
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ── Column specification ──────────────────────────────────────────────────────
# (header_name, json_key, column_width)
# IMPORTANT: json_key must exactly match what Claude returns.
# Col 13 key has trailing space to match the prompt spec JSON key.

COLUMNS = [
    ("S. No",                                    "S. No",                                     8),
    ("Difficulty",                               "Difficulty",                               12),
    ("Question Content (in Markdown)",           "Question Content (in Markdown)",            85),
    ("Question_short_text",                      "Question_short_text",                      25),
    ("Sub_Topics",                               "Sub_Topics",                               18),
    ("Code_language",                            "Code_language",                            15),
    ("Solution Code",                            "Solution Code",                            65),
    ("Front_Code_language",                      "Front_Code_language",                      20),
    ("Prefilled Code For Each Code Language",    "Prefilled Code For Each Code Language",    55),
    ("Test_case_input",                          "Test_case_input",                          30),
    ("Test_case_output",                         "Test_case_output",                         25),
    ("Test_case_type",                           "Test_case_type",                           18),
    ("Code_language",                            "Code_language ",                           15),  # trailing space key
    ("Backend Code",                             "Backend Code",                             25),
]

HEADER_NAMES = [c[0] for c in COLUMNS]
ROW_KEYS     = [c[1] for c in COLUMNS]
COL_WIDTHS   = [c[2] for c in COLUMNS]

# ── Style constants ───────────────────────────────────────────────────────────
HEADER_FONT  = Font(name="Arial", size=10, bold=True, color="FFFFFF")
DATA_FONT    = Font(name="Arial", size=10)
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
WRAP_TOP     = Alignment(wrap_text=True, vertical="top")
CENTER_MID   = Alignment(horizontal="center", vertical="center", wrap_text=True)

_side        = Side(style="thin", color="BFBFBF")
THIN_BORDER  = Border(left=_side, right=_side, top=_side, bottom=_side)


def _write_header(ws) -> None:
    for col_idx, name in enumerate(HEADER_NAMES, start=1):
        cell            = ws.cell(row=1, column=col_idx, value=name)
        cell.font       = HEADER_FONT
        cell.fill       = HEADER_FILL
        cell.alignment  = CENTER_MID
        cell.border     = THIN_BORDER
    ws.row_dimensions[1].height = 28


def _set_col_widths(ws) -> None:
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_data_row(ws, excel_row: int, row_dict: dict) -> None:
    max_lines = 1
    for col_idx, key in enumerate(ROW_KEYS, start=1):
        val             = row_dict.get(key, "") or ""
        cell            = ws.cell(row=excel_row, column=col_idx, value=val)
        cell.font       = DATA_FONT
        cell.alignment  = WRAP_TOP
        cell.border     = THIN_BORDER
        if isinstance(val, str) and val:
            max_lines = max(max_lines, val.count("\n") + 1)
    # One line ~14pt tall; cap so a huge stress-test array doesn't blow up the row.
    ws.row_dimensions[excel_row].height = min(200, max(14, max_lines * 14))


def rows_to_xlsx(rows: list) -> bytes:
    """
    Convert a list of row dicts (from Claude) to xlsx bytes.

    Args:
        rows: List of dicts — one per xlsx row.

    Returns:
        Bytes of the ready-to-download .xlsx file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title        = "BASE_QUE_SET"
    ws.freeze_panes = "A2"   # freeze header row

    _write_header(ws)
    _set_col_widths(ws)

    excel_row = 2
    for row_dict in rows:
        _write_data_row(ws, excel_row, row_dict)
        excel_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()